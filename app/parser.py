"""
文档解析器 - 支持 Word(.docx)、PDF(.pdf)、Excel(.xlsx)、文本(.txt)、Markdown(.md)
每个文档切分成多个 chunk，每个 chunk 携带完整定位信息（"出生证明"）。
"""

import os
import hashlib
from docx import Document


def _parse_doc_via_com(file_path):
    """通过 Word COM 接口解析 .doc 文件（仅 Windows）"""
    import win32com.client
    import pythoncom

    pythoncom.CoInitialize()
    word = None
    doc = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(os.path.abspath(file_path), ReadOnly=True)

        file_hash = compute_file_hash(file_path)
        chunks = []
        title = os.path.basename(file_path)

        for para_idx in range(1, doc.Paragraphs.Count + 1):
            para = doc.Paragraphs(para_idx)
            text = para.Range.Text.strip()
            if not text or len(text) < 2:
                continue
            chunks.append({
                "content": text,
                "file_path": os.path.abspath(file_path),
                "file_hash": file_hash,
                "source_type": "word",
                "title": title,
                "section_path": "",
                "para_index": para_idx - 1,
            })
        return chunks
    finally:
        try:
            if doc:
                doc.Close(False)
        except Exception:
            pass
        try:
            if word:
                word.Quit()
        except Exception:
            pass
        pythoncom.CoUninitialize()


def compute_file_hash(file_path):
    """计算文件 MD5 哈希，用于增量更新判断"""
    h = hashlib.md5()
    with open(file_path, "rb") as f:
        while True:
            chunk = f.read(8192)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def _extract_heading_path(doc, para_idx):
    """提取段落所属的标题路径（如 "第一章 / 1.1 概述"）"""
    headings = []
    for i in range(para_idx - 1, -1, -1):
        if i >= len(doc.paragraphs):
            continue
        p = doc.paragraphs[i]
        if p.style and p.style.name and "Heading" in p.style.name:
            level = 0
            try:
                level = int(p.style.name.replace("Heading", "").strip())
            except ValueError:
                level = 0
            headings.append((level, p.text.strip()))
            if level == 1:
                break
    headings.reverse()
    result = []
    min_level = 99
    for level, text in headings:
        if level <= min_level:
            result.append(text)
            min_level = level
        elif level > min_level:
            result.append(text)
    return " / ".join(result) if result else ""


def parse_docx(file_path):
    """解析 Word 文档，按段落切块"""
    doc = Document(file_path)
    file_hash = compute_file_hash(file_path)
    chunks = []

    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text or len(text) < 2:
            continue

        chunks.append({
            "content": text,
            "file_path": os.path.abspath(file_path),
            "file_hash": file_hash,
            "source_type": "word",
            "title": os.path.basename(file_path),
            "section_path": _extract_heading_path(doc, para_idx),
            "para_index": para_idx,
        })
    return chunks


def parse_pdf(file_path):
    """解析 PDF 文档，按页切块，每页按段落拆分"""
    import pdfplumber
    file_hash = compute_file_hash(file_path)
    chunks = []
    title = os.path.basename(file_path)

    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages):
            text = page.extract_text()
            if not text:
                continue

            # 按空行拆段落
            paragraphs = [p.strip() for p in text.split("\n\n") if p.strip() and len(p.strip()) >= 2]
            if not paragraphs:
                # 没有段落分隔就按行拆
                paragraphs = [p.strip() for p in text.split("\n") if p.strip() and len(p.strip()) >= 2]

            for para_idx, para in enumerate(paragraphs):
                chunks.append({
                    "content": para,
                    "file_path": os.path.abspath(file_path),
                    "file_hash": file_hash,
                    "source_type": "pdf",
                    "title": title,
                    "section_path": f"第 {page_num + 1} 页",
                    "para_index": page_num * 1000 + para_idx,
                })
    return chunks


def parse_excel(file_path):
    """解析 Excel 文档，按 sheet 和行切块"""
    import openpyxl
    file_hash = compute_file_hash(file_path)
    chunks = []
    title = os.path.basename(file_path)

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            # 把行数据拼接成文本
            cells = [str(c) for c in row if c is not None]
            if not cells:
                continue
            row_text = " | ".join(cells)
            if len(row_text.strip()) < 2:
                continue
            chunks.append({
                "content": row_text,
                "file_path": os.path.abspath(file_path),
                "file_hash": file_hash,
                "source_type": "excel",
                "title": title,
                "section_path": f"Sheet: {sheet_name}",
                "para_index": row_idx,
            })
    wb.close()
    return chunks


def parse_text(file_path):
    """解析纯文本/Markdown 文件，按行切块（连续非空行合并）"""
    file_hash = compute_file_hash(file_path)
    chunks = []
    title = os.path.basename(file_path)

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        lines = f.readlines()

    buffer = []
    buffer_start_line = 0
    section_path = ""

    for line_idx, raw_line in enumerate(lines):
        line = raw_line.rstrip("\n").rstrip("\r")
        stripped = line.strip()

        # Markdown 标题跟踪
        if stripped.startswith("#"):
            if buffer:
                chunk_text = "\n".join(buffer)
                if len(chunk_text.strip()) >= 2:
                    chunks.append({
                        "content": chunk_text,
                        "file_path": os.path.abspath(file_path),
                        "file_hash": file_hash,
                        "source_type": "text",
                        "title": title,
                        "section_path": section_path,
                        "para_index": buffer_start_line,
                    })
                buffer = []
            section_path = stripped.lstrip("#").strip()
            buffer_start_line = line_idx

        if stripped:
            buffer.append(line)
        else:
            if buffer:
                chunk_text = "\n".join(buffer)
                if len(chunk_text.strip()) >= 2:
                    chunks.append({
                        "content": chunk_text,
                        "file_path": os.path.abspath(file_path),
                        "file_hash": file_hash,
                        "source_type": "text",
                        "title": title,
                        "section_path": section_path,
                        "para_index": buffer_start_line,
                    })
                buffer = []
            buffer_start_line = line_idx + 1

    if buffer:
        chunk_text = "\n".join(buffer)
        if len(chunk_text.strip()) >= 2:
            chunks.append({
                "content": chunk_text,
                "file_path": os.path.abspath(file_path),
                "file_hash": file_hash,
                "source_type": "text",
                "title": title,
                "section_path": section_path,
                "para_index": buffer_start_line,
            })

    return chunks


def parse_file(file_path):
    """根据扩展名自动选择解析器"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        return parse_docx(file_path)
    elif ext == ".doc":
        return _parse_doc_via_com(file_path)
    elif ext == ".pdf":
        return parse_pdf(file_path)
    elif ext in (".xlsx", ".xls"):
        return parse_excel(file_path)
    elif ext in (".txt", ".md"):
        return parse_text(file_path)
    else:
        return []